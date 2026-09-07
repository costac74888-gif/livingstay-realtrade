"""운영분석 업로드를 저장하지 않고 보수적으로 읽는 파서."""

from __future__ import annotations

import io
import multiprocessing
import os
import re
import signal
import subprocess
import time
import zipfile
from calendar import monthrange
from dataclasses import dataclass
from datetime import date

from openpyxl import load_workbook
from PIL import Image, ImageOps, UnidentifiedImageError


MAX_FILES = 5
MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_TEXT_CHARS = 250_000
MAX_PDF_PAGES = 10
MAX_OCR_PDF_PAGES = 6
MAX_SHEETS = 20
MAX_CELLS = 100_000
MAX_XLSX_UNCOMPRESSED_BYTES = 25 * 1024 * 1024
MAX_IMAGE_PIXELS = 12_000_000
PARSE_TIMEOUT_SECONDS = 45
PARSE_MEMORY_LIMIT_BYTES = 512 * 1024 * 1024
PARSE_SLOT_PATH = "/tmp/livingstay-operation-upload.lock"
ALLOWED_EXTENSIONS = {"csv", "xlsx", "pdf", "jpg", "jpeg", "png"}

_NUMBER = r"(-?\d[\d,\s]*(?:\.\d+)?)"
_DATE = re.compile(
    r"(?<!\d)(20\d{2})[.\-/년]\s*(\d{1,2})(?:[.\-/월]\s*(\d{1,2})일?)?(?!\d)"
)
_PERIOD_LABEL = re.compile(r"(?:분석\s*기간|영업\s*기간|조회\s*기간|기간)\s*[:：]?\s*(.{0,80})", re.I)
_OCC = re.compile(
    rf"(?:\bOCC\b|객실\s*(?:이용|가동)률|객실점유율|가동률)[^0-9-]{{0,20}}{_NUMBER}\s*%?",
    re.I,
)
_ADR = re.compile(
    rf"(?:\bADR\b|판매\s*객실\s*평균\s*요금|평균\s*객실\s*요금)[^0-9-]{{0,20}}{_NUMBER}",
    re.I,
)
_ROOM_REVENUE = re.compile(
    rf"(?:객실\s*매출(?:액)?|룸\s*매출|room\s*revenue)[^0-9-]{{0,20}}{_NUMBER}",
    re.I,
)
_SOLD_ROOMS = re.compile(
    rf"(?:판매\s*객실\s*(?:수|개수|량)|객실\s*판매량|sold\s*rooms?|rooms?\s*sold)[^0-9-]{{0,20}}{_NUMBER}",
    re.I,
)
_HEADER_ALIASES = {
    "period": re.compile(r"^(?:분석|영업|조회)?\s*기간$|^(?:일자|날짜|date|month)$", re.I),
    "occ": re.compile(r"^(?:occ|객실\s*(?:이용|가동)률|객실점유율|가동률)(?:\(%\))?$", re.I),
    "adr": re.compile(r"^(?:adr|판매\s*객실\s*평균\s*요금|평균\s*객실\s*요금)(?:\(원\))?$", re.I),
    "room_revenue": re.compile(r"^(?:객실\s*매출(?:액)?|룸\s*매출|room\s*revenue)(?:\(원\))?$", re.I),
    "sold_rooms": re.compile(r"^(?:판매\s*객실\s*(?:수|개수|량)|객실\s*판매량|sold\s*rooms?|rooms?\s*sold)(?:\(실\))?$", re.I),
}


class DocumentParseError(ValueError):
    """사용자에게 원문이나 파일명을 노출하지 않는 안전한 파싱 오류."""


@dataclass(frozen=True)
class ParsedDocument:
    period_start: str | None = None
    period_end: str | None = None
    occ: float | None = None
    adr: float | None = None
    room_revenue: float | None = None
    sold_rooms: float | None = None
    occupancy_days: int | None = None
    adr_calculation_valid: bool = True

    def as_dict(self):
        adr = self.adr
        adr_source = "explicit" if adr is not None else None
        if (
            adr is None
            and self.adr_calculation_valid
            and self.room_revenue is not None
            and self.sold_rooms
            and self.sold_rooms > 0
        ):
            adr = round(self.room_revenue / self.sold_rooms, 2)
            adr_source = "calculated_from_room_revenue"
        return {
            "period_start": self.period_start,
            "period_end": self.period_end,
            "occ": self.occ,
            "adr": adr,
            "adr_source": adr_source,
            "room_revenue": self.room_revenue,
            "sold_rooms": self.sold_rooms,
            "occupancy_days": self.occupancy_days,
            "adr_calculation_valid": self.adr_calculation_valid,
        }


def _number(value):
    try:
        number = float(re.sub(r"[\s,]", "", value))
    except (TypeError, ValueError):
        return None
    return number if number >= 0 else None


def _first(pattern, text, *, maximum=None):
    match = pattern.search(text)
    value = _number(match.group(1)) if match else None
    if value is not None and maximum is not None and value > maximum:
        return None
    return value


def _iso_date(parts, *, period_end=False):
    year, month, day = parts
    year = int(year)
    if year < 100:
        year += 2000
    month = int(month)
    if not (2000 <= year <= 2100 and 1 <= month <= 12):
        return None
    day = int(day or (monthrange(year, month)[1] if period_end else 1))
    if not (2000 <= year <= 2100 and 1 <= month <= 12 and 1 <= day <= 31):
        return None
    return f"{year:04d}-{month:02d}-{day:02d}"


def _period(text):
    labelled = _PERIOD_LABEL.search(text)
    candidates = list(_DATE.finditer(labelled.group(1) if labelled else text))
    if not candidates:
        return None, None
    start = _iso_date(candidates[0].groups())
    end_match = candidates[1] if len(candidates) > 1 else candidates[0]
    end = _iso_date(
        end_match.groups(),
        period_end=end_match.group(3) is None,
    )
    return (start, end) if start and end else (None, None)


def _inclusive_days(period_start, period_end):
    try:
        start = date.fromisoformat(str(period_start))
        end = date.fromisoformat(str(period_end))
    except (TypeError, ValueError):
        return None
    days = (end - start).days + 1
    return days if 0 < days <= 3660 else None


def parse_metrics(text):
    normalized = re.sub(r"[ \t]+", " ", str(text or ""))[:MAX_TEXT_CHARS]
    period_start, period_end = _period(normalized)
    sold_rooms = _first(_SOLD_ROOMS, normalized)
    return ParsedDocument(
        period_start=period_start,
        period_end=period_end,
        occ=_first(_OCC, normalized, maximum=100),
        adr=_first(_ADR, normalized),
        room_revenue=_first(_ROOM_REVENUE, normalized),
        sold_rooms=sold_rooms,
        occupancy_days=(
            _inclusive_days(period_start, period_end)
            if sold_rooms is not None else None
        ),
    )


def _header_kind(value):
    text = re.sub(r"[\s:：|]+", " ", str(value or "")).strip()
    return next((kind for kind, pattern in _HEADER_ALIASES.items() if pattern.fullmatch(text)), None)


def _cell_number(value, *, maximum=None):
    match = re.search(_NUMBER, str(value or ""))
    number = _number(match.group(1)) if match else None
    if number is not None and maximum is not None and number > maximum:
        return None
    return number


def _cell_period(value):
    matches = list(_DATE.finditer(str(value or "")))
    if not matches:
        return None, None
    start = _iso_date(matches[0].groups())
    end_match = matches[1] if len(matches) > 1 else matches[0]
    end = _iso_date(
        end_match.groups(),
        period_end=end_match.group(3) is None,
    )
    return (start, end) if start and end else (None, None)


def _safe_occ(records):
    """분모 없이도 통합값이 달라질 수 없는 OCC만 반환한다."""
    occ_records = [record for record in records if record.get("occ") is not None]
    if not occ_records:
        return None
    values = {round(float(record["occ"]), 8) for record in occ_records}
    if len(occ_records) > 1:
        return occ_records[0]["occ"] if len(occ_records) == len(records) and len(values) == 1 else None
    if len(records) == 1:
        return occ_records[0]["occ"]
    occ_period = (
        occ_records[0].get("period_start"),
        occ_records[0].get("period_end"),
    )
    periods = [
        (record.get("period_start"), record.get("period_end"))
        for record in records
    ]
    if all(start and end for start, end in periods) and set(periods) == {occ_period}:
        return occ_records[0]["occ"]
    return None


def _safe_explicit_adr(records):
    adr_records = [record for record in records if record.get("adr") is not None]
    if not adr_records:
        return None
    values = {round(float(record["adr"]), 8) for record in adr_records}
    if len(adr_records) > 1:
        return adr_records[0]["adr"] if len(adr_records) == len(records) and len(values) == 1 else None
    if len(records) == 1:
        return adr_records[0]["adr"]
    adr_period = (
        adr_records[0].get("period_start"),
        adr_records[0].get("period_end"),
    )
    periods = [
        (record.get("period_start"), record.get("period_end"))
        for record in records
    ]
    if all(start and end for start, end in periods) and set(periods) == {adr_period}:
        return adr_records[0]["adr"]
    return None


def _occupancy_days(records):
    intervals = []
    for record in records:
        if record.get("sold_rooms") is None:
            continue
        try:
            start = date.fromisoformat(str(record.get("period_start")))
            end = date.fromisoformat(str(record.get("period_end")))
        except (TypeError, ValueError):
            return None
        if end < start:
            return None
        intervals.append((start, end))
    if not intervals:
        return None
    intervals.sort()
    total = 0
    previous_end = None
    for start, end in intervals:
        if previous_end is not None and start <= previous_end:
            return None
        total += (end - start).days + 1
        previous_end = end
    return total if 0 < total <= 3660 else None


def parse_tabular_rows(rows):
    rows = [[str(value).strip() if value is not None else "" for value in row] for row in rows]
    rows = [row for row in rows if any(row)]
    for header_index, row in enumerate(rows):
        columns = {kind: index for index, value in enumerate(row) if (kind := _header_kind(value))}
        metric_columns = set(columns) & {"occ", "adr", "room_revenue", "sold_rooms"}
        if len(metric_columns) < 2 and not ("period" in columns and metric_columns):
            continue
        records = []
        for data_row in rows[header_index + 1:]:
            if all(not str(data_row[index]).strip() for index in columns.values() if index < len(data_row)):
                continue
            record = {}
            for kind, index in columns.items():
                value = data_row[index] if index < len(data_row) else ""
                if kind == "period":
                    record["period_start"], record["period_end"] = _cell_period(value)
                else:
                    record[kind] = _cell_number(value, maximum=100 if kind == "occ" else None)
            if any(record.get(kind) is not None for kind in metric_columns):
                records.append(record)
        if not records:
            break
        starts = [record["period_start"] for record in records if record.get("period_start")]
        ends = [record["period_end"] for record in records if record.get("period_end")]
        revenue_records = [
            record for record in records
            if record.get("room_revenue") is not None
        ]
        sold_records = [
            record for record in records
            if record.get("sold_rooms") is not None and record["sold_rooms"] > 0
        ]
        paired = [
            record for record in records
            if record.get("room_revenue") is not None
            and record.get("sold_rooms") is not None
            and record["sold_rooms"] > 0
        ]
        pairing_valid = (
            bool(paired)
            and len(paired) == len(revenue_records)
            and len(paired) == len(sold_records)
        )
        pairing_eligible = pairing_valid or not (revenue_records and sold_records)
        revenues = [record["room_revenue"] for record in revenue_records]
        sold = [record["sold_rooms"] for record in sold_records]
        occ = _safe_occ(records)
        adr = None if pairing_valid else _safe_explicit_adr(records)
        return ParsedDocument(
            period_start=min(starts) if starts else None,
            period_end=max(ends) if ends else None,
            occ=round(occ, 2) if occ is not None else None,
            adr=round(adr, 2) if adr is not None else None,
            room_revenue=sum(revenues) if revenues else None,
            sold_rooms=sum(sold) if sold else None,
            occupancy_days=_occupancy_days(sold_records),
            adr_calculation_valid=pairing_eligible,
        )

    # 세로형 요약표는 각 행의 첫 셀을 라벨, 나머지를 하나의 값으로 읽는다.
    summary_lines = []
    for row in rows:
        if not row:
            continue
        label = row[0]
        value = "".join(row[1:])
        summary_lines.append(f"{label}: {value}")
    return parse_metrics("\n".join(summary_lines))


def _csv_rows(raw):
    for encoding in ("utf-8-sig", "cp949", "euc-kr"):
        try:
            decoded = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise DocumentParseError("CSV 문자 인코딩을 읽을 수 없습니다.")
    import csv
    return list(csv.reader(io.StringIO(decoded)))[:MAX_CELLS]


def _xlsx_rows(raw):
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if sum(info.file_size for info in archive.infolist()) > MAX_XLSX_UNCOMPRESSED_BYTES:
                raise DocumentParseError("엑셀 내용이 너무 큽니다.")
    except (zipfile.BadZipFile, OSError) as exc:
        raise DocumentParseError("엑셀 내용을 읽을 수 없습니다.") from exc
    try:
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:
        raise DocumentParseError("엑셀 내용을 읽을 수 없습니다.") from exc
    output, cell_count = [], 0
    try:
        for sheet in workbook.worksheets[:MAX_SHEETS]:
            for row in sheet.iter_rows(values_only=True):
                values = list(row)
                if any(value is not None for value in values):
                    output.append(values)
                cell_count += len(row)
                if cell_count >= MAX_CELLS:
                    return output
    finally:
        workbook.close()
    return output


def _remaining_timeout(deadline, maximum):
    remaining = deadline - time.monotonic()
    if remaining <= 0:
        raise DocumentParseError("자료 분석 시간이 초과되었습니다.")
    return max(1, min(maximum, int(remaining)))


def _ocr_image(image, deadline):
    if image.width * image.height > MAX_IMAGE_PIXELS:
        raise DocumentParseError("이미지 해상도가 너무 큽니다.")
    image = ImageOps.exif_transpose(image).convert("L")
    image.thumbnail((3000, 3000))
    buffer = io.BytesIO()
    image.save(buffer, format="PNG", optimize=True)
    try:
        result = subprocess.run(
            ["tesseract", "stdin", "stdout", "-l", "kor+eng", "--psm", "6"],
            input=buffer.getvalue(), stdout=subprocess.PIPE, stderr=subprocess.DEVNULL,
            timeout=_remaining_timeout(deadline, 15), check=False,
        )
    except (OSError, subprocess.TimeoutExpired) as exc:
        raise DocumentParseError("이미지 글자를 읽을 수 없습니다.") from exc
    if result.returncode:
        raise DocumentParseError("이미지 글자를 읽을 수 없습니다.")
    return result.stdout.decode("utf-8", errors="replace")


def _image_text(raw, deadline):
    try:
        with Image.open(io.BytesIO(raw)) as image:
            if image.width * image.height > MAX_IMAGE_PIXELS:
                raise DocumentParseError("이미지 해상도가 너무 큽니다.")
            image.verify()
        with Image.open(io.BytesIO(raw)) as image:
            return _ocr_image(image, deadline)
    except (UnidentifiedImageError, OSError, Image.DecompressionBombError) as exc:
        raise DocumentParseError("이미지 내용을 읽을 수 없습니다.") from exc


def _pdf_text(raw, deadline):
    try:
        import pymupdf
        document = pymupdf.open(stream=raw, filetype="pdf")
    except Exception as exc:
        raise DocumentParseError("PDF 내용을 읽을 수 없습니다.") from exc
    output = []
    try:
        if document.page_count > MAX_PDF_PAGES:
            raise DocumentParseError(f"PDF는 {MAX_PDF_PAGES}쪽 이하만 분석할 수 있습니다.")
        ocr_pages = 0
        for page in document:
            _remaining_timeout(deadline, PARSE_TIMEOUT_SECONDS)
            text = page.get_text("text")[:MAX_TEXT_CHARS]
            output.append(text)
            if len(text.strip()) < 20:
                ocr_pages += 1
                if ocr_pages > MAX_OCR_PDF_PAGES:
                    raise DocumentParseError(f"글자 인식이 필요한 PDF는 {MAX_OCR_PDF_PAGES}쪽 이하만 분석할 수 있습니다.")
                rect = page.rect
                scale = min(2.0, 2200 / max(float(rect.width), 1), 2200 / max(float(rect.height), 1))
                pixmap = page.get_pixmap(matrix=pymupdf.Matrix(scale, scale), alpha=False)
                with Image.open(io.BytesIO(pixmap.tobytes("png"))) as image:
                    output.append(_ocr_image(image, deadline))
            if sum(len(part) for part in output) >= MAX_TEXT_CHARS:
                break
    finally:
        document.close()
    return "\n".join(output)


def parse_document(raw, extension, *, deadline=None):
    deadline = deadline or (time.monotonic() + PARSE_TIMEOUT_SECONDS)
    extension = str(extension or "").lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise DocumentParseError("지원하지 않는 파일 형식입니다.")
    if not raw or len(raw) > MAX_FILE_BYTES:
        raise DocumentParseError("파일은 5MB 이하이어야 합니다.")
    if extension == "csv":
        return parse_tabular_rows(_csv_rows(raw)).as_dict()
    elif extension == "xlsx":
        return parse_tabular_rows(_xlsx_rows(raw)).as_dict()
    elif extension == "pdf":
        if not raw.startswith(b"%PDF"):
            raise DocumentParseError("PDF 내용을 읽을 수 없습니다.")
        text = _pdf_text(raw, deadline)
    else:
        text = _image_text(raw, deadline)
    return parse_metrics(text).as_dict()


def _parse_batch_worker(connection, items, memory_limit_bytes, cpu_limit_seconds):
    """격리된 자식 프로세스에서만 실행되며 원문/파일명을 부모 오류로 보내지 않는다."""
    try:
        try:
            import resource
            resource.setrlimit(resource.RLIMIT_AS, (memory_limit_bytes, memory_limit_bytes))
            resource.setrlimit(resource.RLIMIT_CPU, (cpu_limit_seconds, cpu_limit_seconds + 1))
        except (ImportError, OSError, ValueError):
            pass
        deadline = time.monotonic() + cpu_limit_seconds
        documents = [
            parse_document(raw, extension, deadline=deadline)
            for raw, extension in items
        ]
        connection.send(("ok", documents))
    except DocumentParseError as exc:
        connection.send(("user_error", str(exc)))
    except BaseException:
        connection.send(("internal_error", None))
    finally:
        connection.close()


def _isolated_entry(connection, items, memory_limit_bytes, cpu_limit_seconds, worker_target):
    try:
        os.setsid()
    except OSError:
        pass
    worker_target(connection, items, memory_limit_bytes, cpu_limit_seconds)


def _stop_process_group(process):
    if not process.is_alive():
        return
    try:
        os.killpg(process.pid, signal.SIGTERM)
    except (ProcessLookupError, PermissionError):
        process.terminate()
    process.join(2)
    if process.is_alive():
        try:
            os.killpg(process.pid, signal.SIGKILL)
        except (ProcessLookupError, PermissionError):
            process.kill()
        process.join(2)


def parse_documents_isolated(items, *, timeout_seconds=PARSE_TIMEOUT_SECONDS, worker_target=None):
    """하드 wall-clock 제한이 있는 별도 프로세스에서 문서 묶음을 분석한다."""
    context = multiprocessing.get_context("spawn")
    parent, child = context.Pipe(duplex=False)
    process = context.Process(
        target=_isolated_entry,
        args=(
            child, items, PARSE_MEMORY_LIMIT_BYTES, max(1, int(timeout_seconds)),
            worker_target or _parse_batch_worker,
        ),
        daemon=True,
    )
    process.start()
    child.close()
    try:
        if not parent.poll(timeout_seconds):
            _stop_process_group(process)
            raise DocumentParseError("자료 분석 시간이 초과되었습니다.")
        status, payload = parent.recv()
        process.join(2)
        if status == "ok":
            return payload
        if status == "user_error":
            raise DocumentParseError(payload)
        raise DocumentParseError("자료를 분석하지 못했습니다. 파일 형식과 내용을 확인해 주세요.")
    except EOFError as exc:
        raise DocumentParseError("자료를 분석하지 못했습니다. 파일 형식과 내용을 확인해 주세요.") from exc
    finally:
        parent.close()
        _stop_process_group(process)


def try_acquire_parse_slot():
    """Gunicorn 워커 전체에서 운영자료 분석 한 건만 허용한다."""
    import fcntl
    handle = open(PARSE_SLOT_PATH, "a+b")
    try:
        fcntl.flock(handle.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
        return handle
    except BlockingIOError:
        handle.close()
        return None


def release_parse_slot(handle):
    if not handle:
        return
    import fcntl
    try:
        fcntl.flock(handle.fileno(), fcntl.LOCK_UN)
    finally:
        handle.close()


def merge_documents(documents):
    found = [item for item in documents if any(
        item.get(key) is not None
        for key in ("period_start", "occ", "adr", "room_revenue", "sold_rooms")
    )]
    if not found:
        return ParsedDocument().as_dict()
    starts = [item["period_start"] for item in found if item.get("period_start")]
    ends = [item["period_end"] for item in found if item.get("period_end")]
    revenue_documents = [
        item for item in found if item.get("room_revenue") is not None
    ]
    sold_documents = [
        item for item in found
        if item.get("sold_rooms") is not None and item["sold_rooms"] > 0
    ]
    room_revenue = (
        revenue_documents[0]["room_revenue"]
        if len(revenue_documents) == 1 else None
    )
    sold_rooms = (
        sold_documents[0]["sold_rooms"]
        if len(sold_documents) == 1 else None
    )
    revenue_period = (
        revenue_documents[0].get("period_start"),
        revenue_documents[0].get("period_end"),
    ) if len(revenue_documents) == 1 else (None, None)
    sold_period = (
        sold_documents[0].get("period_start"),
        sold_documents[0].get("period_end"),
    ) if len(sold_documents) == 1 else (None, None)
    compatible_pairing = (
        room_revenue is not None
        and sold_rooms is not None
        and all(revenue_period)
        and revenue_period == sold_period
        and revenue_documents[0].get("adr_calculation_valid", True)
        and sold_documents[0].get("adr_calculation_valid", True)
    )
    occ = _safe_occ(found)
    explicit_adr_records = [
        item for item in found
        if item.get("adr") is not None and item.get("adr_source") == "explicit"
    ]
    explicit_adr = _safe_explicit_adr(explicit_adr_records)
    found_periods = {
        (item.get("period_start"), item.get("period_end")) for item in found
    }
    if explicit_adr is not None and (
        (None, None) in found_periods or len(found_periods) > 1
    ):
        explicit_adr = None
    return ParsedDocument(
        period_start=min(starts) if starts else None,
        period_end=max(ends) if ends else None,
        occ=round(occ, 2) if occ is not None else None,
        adr=None if compatible_pairing else explicit_adr,
        room_revenue=room_revenue,
        sold_rooms=sold_rooms,
        occupancy_days=(
            sold_documents[0].get("occupancy_days")
            if len(sold_documents) == 1 else None
        ),
        adr_calculation_valid=compatible_pairing,
    ).as_dict()