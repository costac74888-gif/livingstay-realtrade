"""Approved annual tourism-accommodation roster ingestion and statistics.

This is intentionally independent from the Tourism Data Lab (CSV) pipeline and
from ``lodging_registry``.  The annual Ministry workbook is an aggregate
evidence source, not a building-master or permit-registry replacement.
"""
from __future__ import annotations

import hashlib
import io
import json
import os
import re
import secrets
from datetime import datetime
from itertools import chain

from openpyxl import load_workbook
from psycopg2.extras import execute_values
from addr_norm import normalize_jibun_prefix, normalize_road_prefix

MAX_FILE_BYTES = 25 * 1024**2
MAX_ROWS = 100_000
MAX_SOURCE_NAME_LENGTH = 200
_STAGE_TTL_SECONDS = 600
_YEAR_RE = re.compile(r"(?<!\d)((?:20)?\d{2})(?:년|[-./]\s*12|$)")
_SUBTYPE_MAP = {
    "수상관광호텔": "수상관광호텔업", "의료관광호텔": "의료관광호텔업",
    "관광호텔": "관광호텔업", "관광호텔업": "관광호텔업",
    "휴양콘도": "휴양콘도미니엄업", "휴양콘도미니엄": "휴양콘도미니엄업",
    "휴양콘도미니엄업": "휴양콘도미니엄업", "호스텔": "호스텔업",
    "소형호텔": "소형호텔업", "한국전통호텔": "한국전통호텔업",
    "가족호텔": "가족호텔업",
}


def _actor(value):
    try:
        value = int(value)
    except (TypeError, ValueError):
        raise ValueError("유효한 관리자 정보가 필요합니다.") from None
    if value <= 0:
        raise ValueError("유효한 관리자 정보가 필요합니다.")
    return value


def _safe_filename(name):
    return bool(name and os.path.basename(name) == name and
                re.fullmatch(r"[^/\\\x00-\x1f]{1,180}\.xlsx", name, re.I))


def _text(value):
    return str(value or "").strip()


def _year(value, label):
    """Accept only an explicit four-digit annual-roster year."""
    raw = _text(value)
    if not re.fullmatch(r"\d{4}", raw):
        raise ValueError(f"{label}은 2000~2100년의 네 자리 숫자로 입력해야 합니다.")
    year = int(raw)
    if not 2000 <= year <= 2100:
        raise ValueError(f"{label}은 2000~2100년 범위여야 합니다.")
    return year


def validate_metadata(reference_year, next_collection_year, source_name, workbook_year):
    """Validate human-entered collection metadata against the parsed workbook."""
    reference_year = _year(reference_year, "기준 연도")
    next_collection_year = _year(next_collection_year, "다음 수집 연도")
    source_name = _text(source_name)
    if not source_name or len(source_name) > MAX_SOURCE_NAME_LENGTH:
        raise ValueError(
            f"수집 출처명은 1~{MAX_SOURCE_NAME_LENGTH}자의 비어 있지 않은 텍스트여야 합니다."
        )
    if next_collection_year < reference_year:
        raise ValueError("다음 수집 연도는 기준 연도보다 같거나 커야 합니다.")
    if reference_year != workbook_year:
        raise ValueError("입력한 기준 연도가 XLSX 파일명에서 확인한 기준 연도와 다릅니다.")
    return {
        "reference_year": reference_year,
        "next_collection_year": next_collection_year,
        "source_name": source_name,
    }


def _integer(value):
    if value is None or _text(value) in ("", "-", "N/A"):
        return 0
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return max(0, int(value))
    cleaned = re.sub(r"[^0-9-]", "", _text(value).replace(",", ""))
    try:
        return max(0, int(cleaned))
    except ValueError:
        return 0


def _subtype(value):
    value = re.sub(r"\s+", "", _text(value))
    for token, normalized in _SUBTYPE_MAP.items():
        if token in value:
            return normalized
    return value or "미분류"


def _address_key(value):
    return normalize_road_prefix(value) or normalize_jibun_prefix(value) or None


def _header_key(value):
    return re.sub(r"[\s\n\r()（）]", "", _text(value)).lower()


def _column(headers, *terms):
    for index, header in enumerate(headers):
        key = _header_key(header)
        if any(term in key for term in terms):
            return index
    return None


def parse_xlsx(filename, raw):
    """Parse the published annual workbook without trusting formulas/macros."""
    if not _safe_filename(filename) or not raw or len(raw) > MAX_FILE_BYTES:
        raise ValueError("XLSX 파일명 또는 크기가 허용 범위를 벗어났습니다.")
    if not raw.startswith(b"PK\x03\x04"):
        raise ValueError("유효한 XLSX 파일이 아닙니다.")
    try:
        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True,
                             keep_links=False)
    except Exception as exc:
        raise ValueError("XLSX를 안전하게 읽을 수 없습니다.") from exc
    if len(book.sheetnames) != 1:
        raise ValueError("연간 관광숙박 명부는 시트 하나만 허용됩니다.")
    sheet = book[book.sheetnames[0]]
    rows = sheet.iter_rows(values_only=True)
    # The official fixture uses merged, two-level headings.  Locate the
    # physical header row rather than assuming row 1, while retaining only
    # one predictable data boundary.
    leading = [list(row) for _, row in zip(range(10), rows)]
    header_at = next((index for index, row in enumerate(leading)
                      if _column([_text(v) for v in row], "업종") is not None), None)
    if header_at is None:
        raise ValueError("필수 헤더(업종, 영업상태, 객실수)를 찾을 수 없습니다.")
    header_end = next((index for index in range(header_at, min(len(leading), header_at + 3))
                       if _column([_text(v) for v in leading[index]], "객실") is not None), None)
    if header_end is None:
        raise ValueError("필수 헤더(업종, 영업상태, 객실수)를 찾을 수 없습니다.")
    headers = [
        " ".join(_text(leading[level][column]) if column < len(leading[level]) else ""
                 for level in range(header_at, header_end + 1)
                 if column < len(leading[level]) and _text(leading[level][column]))
        for column in range(max(len(row) for row in leading[header_at:header_end + 1]))
    ]
    subtype_col = _column(headers, "업종", "관광숙박업종")
    region1_col = _column(headers, "지역1", "시도")
    region2_col = _column(headers, "지역2", "시군구")
    facility_col = _column(headers, "시설개요", "시설명", "관광사업자명")
    address_col = _column(headers, "주소")
    status_col = _column(headers, "영업상태")
    rooms_col = _column(headers, "객실수", "객실")
    if subtype_col is None or rooms_col is None or status_col is None:
        raise ValueError("필수 헤더(업종, 객실수, 영업상태)를 찾을 수 없습니다.")
    parsed = []
    data_rows = chain(leading[header_end + 1:], rows)
    for row_number, values in enumerate(data_rows, header_end + 2):
        if row_number > MAX_ROWS + header_end + 1:
            raise ValueError("XLSX 행 수 한도를 초과했습니다.")
        values = list(values)
        if not any(v not in (None, "") for v in values):
            continue
        def at(index):
            return _text(values[index]) if index is not None and index < len(values) else ""
        raw_subtype = at(subtype_col)
        if not raw_subtype:
            continue
        raw_status = at(status_col)
        parsed.append({
            "row_number": row_number, "sido_name": at(region1_col),
            "sgg_name": at(region2_col), "facility_name": at(facility_col),
            "address": at(address_col), "address_norm": _address_key(at(address_col)) or "",
            "subtype": _subtype(raw_subtype), "raw_subtype": raw_subtype,
            "room_count": _integer(at(rooms_col)), "raw_status": raw_status,
            # Published roster policy is intentionally exact: 휴업/영업종료/
            # blank values are not silently included in active metrics.
            "is_active": raw_status == "영업중",
            "status_bucket": "active" if raw_status == "영업중" else (
                "inactive" if raw_status else "review"
            ),
        })
    if not parsed:
        raise ValueError("유효한 관광숙박 명부 행이 없습니다.")
    year_match = _YEAR_RE.search(filename)
    if not year_match:
        raise ValueError("파일명에서 기준 연도를 확인할 수 없습니다.")
    year = int(year_match.group(1))
    return (year + 2000 if year < 100 else year), parsed, headers


def preview(file, owner, reference_year, next_collection_year, source_name):
    owner = _actor(owner)
    if file is None:
        raise ValueError("XLSX 파일이 필요합니다.")
    filename, raw = file.filename or "", file.read()
    year, rows, headers = parse_xlsx(filename, raw)
    metadata = validate_metadata(
        reference_year, next_collection_year, source_name, year
    )
    digest = hashlib.sha256(raw).hexdigest()
    token = secrets.token_urlsafe(32)
    manifest = {"year": year, **metadata, "source_file": filename, "sha256": digest,
                "headers": headers, "rows": rows}
    active = [row for row in rows if row["is_active"]]
    return token, owner, manifest, {
        **metadata, "source_file": filename, "sha256": digest,
        "total_rows": len(rows), "active_facilities": len(active),
        "active_facility_count": len(active),
        "active_rooms": sum(row["room_count"] for row in active),
        "active_room_count": sum(row["room_count"] for row in active),
        "inactive_count": sum(row["status_bucket"] == "inactive" for row in rows),
        "review_count": sum(row["status_bucket"] == "review" for row in rows),
        "inactive_or_review_count": len(rows) - len(active),
        "permit_count": len(active), "room_count": sum(row["room_count"] for row in active),
        "subtypes": _breakdown(active),
    }


def _breakdown(rows):
    result = {}
    for row in rows:
        item = result.setdefault(row["subtype"], {"subtype": row["subtype"], "permit_count": 0, "room_count": 0})
        item["permit_count"] += 1
        item["room_count"] += int(row["room_count"] or 0)
    return sorted(result.values(), key=lambda item: item["subtype"])


def cross_check_rows(conn, rows):
    """Conservative, reproducible master comparison for preview/apply evidence."""
    cur = conn.cursor()
    try:
        cur.execute("SELECT id, building_name, road_address, jibun_address FROM master_buildings")
        by_address = {}
        for building in cur.fetchall():
            name = re.sub(r"[^0-9A-Za-z가-힣]", "", _text(building["building_name"])).lower()
            for address in (building["road_address"], building["jibun_address"]):
                key = _address_key(address)
                if name and key:
                    by_address.setdefault(key, {}).setdefault(name, set()).add(building["id"])
    finally:
        cur.close()
    counts = {"matched": 0, "unmatched": 0, "ambiguous": 0, "conflict": 0}
    subtype_ids = {row["subtype"]: set() for row in rows}
    checked = []
    for source in rows:
        row = dict(source)
        name = re.sub(r"[^0-9A-Za-z가-힣]", "", _text(row.get("facility_name"))).lower()
        candidates_by_name = by_address.get(row.get("address_norm") or "", {})
        matches = candidates_by_name.get(name, set()) if name else set()
        if len(matches) == 1:
            status, building_id = "matched", next(iter(matches))
            subtype_ids.setdefault(row["subtype"], set()).add(building_id)
        elif len(matches) > 1:
            status, building_id = "ambiguous", None
        elif name and row.get("address_norm") and candidates_by_name:
            # Same verified address but a different master name: never guess.
            status, building_id = "conflict", None
        else:
            status, building_id = "unmatched", None
        row["cross_check_status"] = status
        row["matched_building_id"] = building_id
        counts[status] += 1
        checked.append(row)
    subtype_linked = [
        {"subtype": subtype, "linked_building_count": len(ids)}
        for subtype, ids in sorted(subtype_ids.items())
    ]
    all_linked = set().union(*subtype_ids.values()) if subtype_ids else set()
    return checked, {
        **counts, "linked_building_count": len(all_linked),
        "subtype_linked_buildings": subtype_linked,
    }


def store_preview(conn, file, owner, reference_year, next_collection_year, source_name):
    token, owner, manifest, summary = preview(
        file, owner, reference_year, next_collection_year, source_name
    )
    checked_rows, evidence = cross_check_rows(conn, manifest["rows"])
    manifest["rows"] = checked_rows
    manifest["cross_check"] = evidence
    cur = conn.cursor()
    try:
        cur.execute("""INSERT INTO annual_tourism_roster_stages
          (token, admin_user_id, manifest, expires_at, state)
          VALUES (%s,%s,%s::jsonb,NOW() + interval '10 minutes','previewed')""",
          (token, owner, json.dumps(manifest, ensure_ascii=False)))
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        cur.close()
    blocking_conflicts = evidence["ambiguous"] + evidence["conflict"]
    return {"token": token, "expires_in_seconds": _STAGE_TTL_SECONDS,
            "safe": blocking_conflicts == 0,
            "blocking_conflicts": blocking_conflicts,
            "building_cross_check": evidence, **summary}


def assert_production_connection(conn):
    """Annual approved-source writes are deliberately production-only."""
    # Replit sets this only inside a published app. Comparing DATABASE_URL with
    # a separately stored PROD_DATABASE_URL is unreliable because the same
    # production database can be reached through different proxy endpoints.
    if os.environ.get("REPLIT_DEPLOYMENT") != "1":
        raise RuntimeError("승인 연간 관광숙박 명부는 운영 서버에서만 적용할 수 있습니다.")


def apply(conn, token, owner):
    owner = _actor(owner)
    assert_production_connection(conn)
    cur = conn.cursor()
    try:
        cur.execute("""UPDATE annual_tourism_roster_stages SET state='applying'
          WHERE token=%s AND admin_user_id=%s AND state='previewed' AND expires_at > NOW()
          RETURNING manifest""", (token, owner))
        stage = cur.fetchone()
        if not stage:
            raise ValueError("미리보기 토큰이 없거나 이미 적용·만료되었습니다.")
        manifest = stage["manifest"]
        year, source_file, digest = manifest["year"], manifest["source_file"], manifest["sha256"]
        metadata = validate_metadata(
            manifest.get("reference_year"),
            manifest.get("next_collection_year"),
            manifest.get("source_name"),
            year,
        )
        # Do this before creating/changing a version.  The check is repeated
        # even though preview recorded it, because the master may have changed.
        checked_rows, check = cross_check_rows(conn, manifest["rows"])
        if check["ambiguous"] or check["conflict"]:
            raise ValueError("건물 대조에 모호 또는 충돌 행이 있어 적용할 수 없습니다.")
        cur.execute("""INSERT INTO annual_tourism_roster_versions
          (reference_year, next_collection_year, source_name, source_file, source_sha256,
           approved_by, approved_at, status)
          VALUES (%s,%s,%s,%s,%s,%s,NOW(),'approved')
          ON CONFLICT (reference_year, source_sha256) DO UPDATE
          SET next_collection_year=EXCLUDED.next_collection_year,
              source_name=EXCLUDED.source_name, approved_by=EXCLUDED.approved_by,
              approved_at=NOW(), status='approved'
          RETURNING id""", (year, metadata["next_collection_year"],
                             metadata["source_name"], source_file, digest, owner))
        version_id = cur.fetchone()["id"]
        cur.execute("DELETE FROM annual_tourism_roster_entries WHERE version_id=%s", (version_id,))
        values = [(version_id, r["row_number"], r["sido_name"], r["sgg_name"],
                   r["facility_name"], r["address"], r["address_norm"], r["subtype"],
                   r["raw_subtype"], r["raw_status"], r["is_active"], r["room_count"])
                  for r in manifest["rows"]]
        execute_values(cur, """INSERT INTO annual_tourism_roster_entries
          (version_id,source_row_number,sido_name,sgg_name,facility_name,address,address_norm,
           subtype,raw_subtype,raw_status,is_active,room_count)
          VALUES %s""", values)
        cur.execute("""SELECT id, source_row_number FROM annual_tourism_roster_entries
                       WHERE version_id=%s""", (version_id,))
        entry_ids = {row["source_row_number"]: row["id"] for row in cur.fetchall()}
        evidence = [
            (entry_ids[row["row_number"]], row["matched_building_id"],
             "name_and_address_exact", row["cross_check_status"])
            for row in checked_rows
        ]
        execute_values(cur, """INSERT INTO annual_tourism_roster_building_evidence
          (entry_id, master_building_id, match_method, match_status) VALUES %s""", evidence)
        cur.execute("""UPDATE annual_tourism_roster_versions SET status='superseded'
                       WHERE reference_year=%s AND id<>%s AND status='approved'""",
                    (year, version_id))
        cur.execute("""UPDATE annual_tourism_roster_stages SET state='applied', applied_at=NOW()
                       WHERE token=%s""", (token,))
        conn.commit()
        return {"version_id": version_id, **metadata, "applied_rows": len(values),
                "building_cross_check": check}
    except Exception:
        conn.rollback()
        raise
    finally:
        cur.close()


def latest_approved_stats(conn):
    """Return only the two metrics that can override the tourism stats row."""
    cur = conn.cursor()
    try:
        cur.execute("""WITH version AS (
              SELECT id, reference_year, next_collection_year, source_name, source_file,
                     source_sha256, approved_at
              FROM annual_tourism_roster_versions WHERE status='approved'
              ORDER BY reference_year DESC, approved_at DESC, id DESC LIMIT 1
            ) SELECT v.reference_year, v.next_collection_year, v.source_name, v.source_file,
                     v.source_sha256, v.approved_at::text,
                     e.subtype, COUNT(*)::int AS permit_count,
                     COALESCE(SUM(e.room_count),0)::int AS room_count,
                     COUNT(DISTINCT x.master_building_id)::int AS linked_building_count
              FROM version v JOIN annual_tourism_roster_entries e ON e.version_id=v.id
              LEFT JOIN annual_tourism_roster_building_evidence x ON x.entry_id=e.id
              WHERE e.is_active
               GROUP BY v.reference_year,v.next_collection_year,v.source_name,v.source_file,
                        v.source_sha256,v.approved_at,e.subtype
              ORDER BY e.subtype""")
        rows = [dict(row) for row in cur.fetchall()]
        cur.execute("""SELECT x.match_status, COUNT(*)::int AS count
          FROM annual_tourism_roster_building_evidence x
          JOIN annual_tourism_roster_entries e ON e.id=x.entry_id
          JOIN annual_tourism_roster_versions v ON v.id=e.version_id
          WHERE v.status='approved'
            AND v.id=(SELECT id FROM annual_tourism_roster_versions WHERE status='approved'
                      ORDER BY reference_year DESC, approved_at DESC, id DESC LIMIT 1)
          GROUP BY x.match_status""")
        evidence = {row["match_status"]: row["count"] for row in cur.fetchall()}
    finally:
        cur.close()
    if not rows:
        return None
    return {"permit_count": sum(row["permit_count"] for row in rows),
            "room_count": sum(row["room_count"] for row in rows),
            "sub_rows": [{"type": row["subtype"], "permit_count": row["permit_count"],
                          "room_count": row["room_count"],
                          "linked_building_count": row["linked_building_count"]} for row in rows],
            "source": {"reference_year": rows[0]["reference_year"],
                        "next_collection_year": rows[0]["next_collection_year"],
                        "source_name": rows[0]["source_name"],
                       "source_file": rows[0]["source_file"],
                       "source_sha256": rows[0]["source_sha256"],
                       "approved_at": rows[0]["approved_at"],
                       "building_cross_check": {"matched": evidence.get("matched", 0),
                                                "unmatched": evidence.get("unmatched", 0)}}}


def latest_linked_subtypes(cur, building_id):
    """Public-safe legal subtype evidence for one conservatively linked building."""
    cur.execute("""SELECT DISTINCT e.subtype
      FROM annual_tourism_roster_building_evidence x
      JOIN annual_tourism_roster_entries e ON e.id=x.entry_id
      JOIN annual_tourism_roster_versions v ON v.id=e.version_id
      WHERE x.master_building_id=%s AND x.match_status='matched' AND e.is_active
        AND v.id=(SELECT id FROM annual_tourism_roster_versions WHERE status='approved'
                  ORDER BY reference_year DESC, approved_at DESC, id DESC LIMIT 1)
      ORDER BY e.subtype""", (building_id,))
    return [row["subtype"] for row in cur.fetchall()]