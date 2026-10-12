"""Build static, privacy-safe review files and an image-embedded Excel/ZIP."""
from pathlib import Path
import csv
import hashlib
import html
import json
import math
import re
import shutil
import zipfile
from PIL import Image as PILImage
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Font, PatternFill

ROOT=Path(__file__).resolve().parents[4]
OUT=ROOT/"generated-artifacts/hs2-full-design"
BUNDLE=OUT/"bundle"

def safe(value):
    return html.escape(str(value))

def main():
    manifest=json.loads((OUT/"screen-manifest.json").read_text())
    results=json.loads((OUT/"capture-report.json").read_text())
    bykey={(v["id"],v["width"]):v for v in results["screens"]}
    screens=manifest["screens"]
    assets=BUNDLE/"assets";assets.mkdir(exist_ok=True)
    for f in (ROOT/"artifacts/mockup-sandbox/public/images/hs2-unified").iterdir():
        if f.suffix.lower() in (".png",".jpg",".jpeg"):
            shutil.copyfile(f,assets/f.name)
    for f in (ROOT/"artifacts/mockup-sandbox/public/images/hs2-complete").iterdir():
        shutil.copyfile(f,assets/f.name)
    base_style="body{margin:0;font:16px/1.6 system-ui,sans-serif;color:#000;background:white}main.report{padding:24px;max-width:1500px;margin:auto}h1{font-size:30px}h2{font-size:24px}a{color:#000}.note{padding:16px;background:#f1f2eb;border:1px solid #596044}table{border-collapse:collapse;width:100%}th,td{text-align:left;padding:12px;border:1px solid #ddd;vertical-align:top}img{max-width:100%}.pair{display:grid;grid-template-columns:minmax(0,1fr) 390px;gap:24px}section{margin:40px 0}nav{display:flex;gap:20px;flex-wrap:wrap}@media(max-width:900px){.pair{display:block}td,th{word-break:break-word}main.report{padding:12px}}"
    rows=[]
    body=[]
    tree={role:[] for role in ("사용자","임대인","관리자")}
    for s in screens:
        web=bykey.get((s["id"],1440),{})
        mobile=bykey.get((s["id"],390),{})
        changes="공통 로고·전체폭 헤더·16px 이상 글자·흰 배경·올리브/겨자 체계; 역할별 정보 구조와 모바일 재배치"
        if s["view"].startswith("search"):
            changes+="; 좌측 사진 2열/우측 지도 44:56; 지역·기간·인원; 불필요한 필터 숨김"
        if s["view"].startswith("detail"):
            changes+="; 갤러리/면적·구조·소개/예약 위계; 주·월 조건; 대략 위치; 모바일 예약바"
        elif s["view"].startswith("registration"):
            changes+="; "+s["title"]+" 전용 단계; 웹 양식 그룹/모바일 단일 열; 이전·다음 로컬 화면 이동"
        elif s["kind"]=="popup":
            changes+="; "+s["title"]+" 독립 열린 상태; 전용 닫기·돌아가기; 모바일 폭 제한"
        elif s["role"]=="관리자":
            changes+="; "+s["title"]+" 전용 항목/목록; 모바일 표 컨테이너 스크롤·양식 재배치"
        elif s["view"].startswith("legacy"):
            changes+="; 원본 제목·항목 기반 양식/목록/분석/설정 유형별 재구성"
        else:
            changes+="; "+s["title"]+" 정보·행동 우선순위 분리; 모바일 카드·버튼 재배치"
        row={"화면ID":s["id"],"역할":s["role"],"상위ID":s["parent"],"화면":s["title"],"종류":s["kind"],
          "원본경로":" | ".join(s["original_paths"]) or "미확인/독립 원본 경로 없음",
          "시안경로":f"/__mockup/preview/hs2-complete/Suite?screen={s['id']}",
          "원본소스":" | ".join(s["source_files"]) or "요청 체크리스트 기반 신규 화면 시안",
          "채택근거":s["basis"],"수정내용":changes,
          "웹판정":web.get("layoutStatus","누락"),"모바일판정":mobile.get("layoutStatus","누락"),
          "기능상태":s["function_status"],"웹캡처":web.get("capture",""),
          "모바일캡처":mobile.get("capture",""),"정적시안":f"screens/{s['id']}.html"}
        rows.append(row);tree[s["role"]].append({"id":s["id"],"parent":s["parent"],"title":s["title"],"kind":s["kind"]})
        for stem in (s["id"],s["id"]+"-web",s["id"]+"-mobile"):
            file=BUNDLE/"screens"/f"{stem}.json"
            if not file.exists():continue
            payload=json.loads(file.read_text())
            css=payload["css"]
            css=re.sub(r'https?://[^\'"\\s)]+(?:appkey|serviceKey)[^\'"\\s)]*','',css,flags=re.I)
            notice='<footer style="padding:18px;background:#f1f2eb;color:#000;font:16px/1.6 sans-serif">다운로드한 정적 화면 시안입니다. 가상 데이터이며 실제 인증·제출·예약·가격 계산·결제는 동작하지 않습니다. 지도는 정적 캡처이며 Noto Sans KR 한글 글꼴과 OFL 라이선스를 함께 제공합니다. <a href="../index.html">비교 목차</a></footer>'
            doc=f'<!doctype html><html lang="ko"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>{safe(s["id"]+" "+s["title"])}</title><style>{css}</style><body>{payload["body"]}{notice}<script>document.addEventListener("submit",function(e){{e.preventDefault();alert("정적 화면 시안에서는 제출되지 않습니다.")}})</script></body></html>'
            (file.with_suffix(".html")).write_text(doc)
            file.unlink()
        # Desktop snapshots and the default responsive file are byte-identical.
        # Keep the web alias without duplicating 246 full HTML documents in the ZIP.
        (BUNDLE/"screens"/f"{s['id']}-web.html").write_text(f'<!doctype html><html lang="ko"><meta charset="utf-8"><meta http-equiv="refresh" content="0;url={s["id"]}.html"><title>{safe(s["title"])}</title><body style="font:16px sans-serif;color:black"><a href="{s["id"]}.html">웹 정적 시안 열기</a></body></html>')
        body.append(f'<section id="{s["id"]}"><h2>{safe(s["id"]+" "+s["title"])}</h2><p>{safe(s["role"])} / {safe(s["kind"])} / 근거: {safe(s["basis"])} / <a href="screens/{s["id"]}.html">웹 정적 시안</a> · <a href="screens/{s["id"]}-mobile.html">모바일 정적 시안</a></p><p>원본: {safe(row["원본경로"])}<br>시안: {safe(row["시안경로"])}</p><p>{safe(row["수정내용"])}<br>웹: {safe(row["웹판정"])} · 모바일: {safe(row["모바일판정"])}<br>{safe(row["기능상태"])}</p><div class="pair"><a href="{safe(row["웹캡처"])}"><img src="{safe(row["웹캡처"])}" alt="{s["id"]} 웹 1440px" loading="lazy"></a><a href="{safe(row["모바일캡처"])}"><img src="{safe(row["모바일캡처"])}" alt="{s["id"]} 모바일 390px" loading="lazy"></a></div></section>')
    nav='<nav>'+''.join(f'<a href="#role-{safe(role)}">{safe(role)}</a>' for role in tree)+'</nav>'
    index='<h2>역할별 화면 트리</h2>'
    for role,nodes in tree.items():
        index+=f'<details id="role-{safe(role)}" open><summary><b>{safe(role)} · {len(nodes)} 화면</b></summary><ul>'
        for n in nodes:index+=f'<li><a href="#{n["id"]}">{safe(n["id"]+" "+n["title"])}</a> / {safe(n["kind"])} / 상위 {safe(n["parent"])}</li>'
        index+='</ul></details>'
    findings={
      "rollback_commit":manifest["rollback_commit"],"screen_count":len(screens),"capture_count":len(results["screens"]),
      "capture_failures":[r["id"] for r in results["screens"] if r.get("error")],
      "design_needs":[{"id":r["id"],"width":r["width"],"metrics":r.get("metrics",{})} for r in results["screens"] if r.get("layoutStatus")!="기본 디자인 검사 통과 · 검토 대기"],
      "unverified_original_routes":[s["id"] for s in screens if not s["original_paths"]],
      "exclusions":manifest["exclusions"],"inventory_limit":manifest["inventory_limit"],
      "functional_pass":False,"external_dependencies":["라이브 지도: Kakao SDK와 기존 개발 Preview의 공개 SDK bootstrap 필요. ZIP에서는 키 없는 지도 캡처로 고정.","글꼴: Google Fonts Noto Sans KR 화면용 부분집합을 로컬 제공. OFL 동봉, 현재 시안 외 새 입력 글리프는 시스템 글꼴로 대체될 수 있음."],
      "not_verified":["경쟁사 인증 내부 화면·전체 소스·코드 재사용 라이선스","운영 인증·실결제·실제 재고·85,617개 시설 연결","로그인 후 서버 데이터로 생성되는 개별 레코드 변형","전체 화면의 최종 사용자 디자인 승인"],
    }
    header=f'<h1>HOME &amp; STAY 전체 화면 비교</h1><p>{len(screens)}개 화면 ID · 웹 1440px / 모바일 390px · 전체 캡처 {len(results["screens"])}장</p>{nav}<p><a href="comparison.xlsx">캡처 삽입 엑셀 다운로드</a> · <a href="screen-status.csv">상태표 CSV</a> · <a href="findings.json">누락·미확인 보고</a></p><div class="note">화면 디자인 검토용이며 기능 PASS가 아닙니다. 모든 데이터는 가상입니다. 운영 인증·실결제·DB 변경·배포를 하지 않았습니다. 시설 실데이터 연결 수량을 주장하지 않습니다. 정적 시안의 지도는 이미지이며 확대/이동·입력 처리·결제는 지원하지 않습니다. 원본 경로는 소스 정의 확인과 운영 실행 검증을 구분합니다.</div><p>롤백 기준: {manifest["rollback_commit"]}</p><p>참고: 삼삼엠투 공개 검색(72px 헤더·263×175px 사진·2열/지도), 에어비앤비 공개 상세(갤러리·정보·예약 패널). 비공개 내부 화면과 전체 소스는 검수하거나 복제하지 않았습니다.</p>'
    (BUNDLE/"index.html").write_text(f'<!doctype html><html lang="ko"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>HOME & STAY 전체 화면 비교</title><style>{base_style}</style><main class="report">{header}{index}{"".join(body)}</main></html>')
    (BUNDLE/"findings.json").write_text(json.dumps(findings,ensure_ascii=False,indent=2))
    (BUNDLE/"role-tree.json").write_text(json.dumps(tree,ensure_ascii=False,indent=2))
    (BUNDLE/"screen-manifest.json").write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
    (BUNDLE/"capture-report.json").write_text(json.dumps(results,ensure_ascii=False,indent=2))
    (BUNDLE/"README.txt").write_text("HOME & STAY 화면 디자인 검토 패키지\n압축을 풀고 index.html을 브라우저로 여세요. comparison.xlsx에는 실제 캡처가 삽입돼 있습니다.\n각 ID-web/mobile.jpg는 전체 화면 캡처입니다. screens/ID.html은 다운로드 가능한 정적 시안입니다.\n기능 완료·권한·결제·운영 자료 연결을 뜻하지 않습니다. 정적 파일은 외부 지도/폰트 요청 없이 캡처 지도와 동봉한 Noto Sans KR 부분집합(OFL)을 사용합니다.\n누락/미확인은 findings.json, 행별 경로/근거/판정은 screen-status.csv를 참고하세요.\n원본 57개 HTML·연결 JS·탭/팝업 정의 기반입니다. 인증 내부·실데이터 변형은 미검수.\n체험/서비스/공동호스트 제외.\n롤백 기준: "+manifest["rollback_commit"])
    with (BUNDLE/"screen-status.csv").open("w",encoding="utf-8-sig",newline="") as f:
        writer=csv.DictWriter(f,fieldnames=list(rows[0]));writer.writeheader();writer.writerows(rows)
    workbook=Workbook();summary=workbook.active;summary.title="화면목록"
    headers=list(rows[0]);summary.append(headers)
    for row in rows:summary.append([row[h] for h in headers])
    summary.freeze_panes="A2";summary.auto_filter.ref=summary.dimensions
    for i,h in enumerate(headers,1):
        summary.cell(1,i).font=Font(bold=True,color="000000")
        summary.cell(1,i).fill=PatternFill("solid",fgColor="E4E8DC")
        summary.column_dimensions[summary.cell(1,i).column_letter].width=18 if i<4 else 40
    for cells in summary.iter_rows(min_row=2):
        for c in cells:c.alignment=Alignment(wrap_text=True,vertical="top")
    for index,s in enumerate(screens,2):
        sheet=workbook.create_sheet(s["id"])
        summary.cell(index,1).hyperlink=f"#'{s['id']}'!A1"
        sheet.merge_cells("A1:P1");sheet["A1"]=s["id"]+" · "+s["title"];sheet["A1"].font=Font(size=18,bold=True)
        sheet.merge_cells("A2:P2");sheet["A2"]="가상 데이터 화면 시안 · 캡처는 기능 PASS가 아님"
        sheet.merge_cells("A3:P3");sheet["A3"]=rows[index-2]["원본경로"]
        sheet.merge_cells("A4:P4");sheet["A4"]=rows[index-2]["시안경로"]
        sheet.merge_cells("A5:H5");sheet["A5"]="웹 1440px · "+rows[index-2]["웹판정"]
        sheet.merge_cells("I5:P5");sheet["I5"]="모바일 390px · "+rows[index-2]["모바일판정"]
        sheet.column_dimensions["A"].width=18
        for letter in "BCDEFGHIJKLMNOP":sheet.column_dimensions[letter].width=12
        largest=0
        for width,col,target in [(1440,"A",740),(390,"J",390)]:
            record=bykey.get((s["id"],width),{})
            file=BUNDLE/record.get("capture","missing")
            if file.is_file():
                im=Image(str(file));ratio=target/im.width;im.width=target;im.height*=ratio;largest=max(largest,im.height)
                sheet.add_image(im,col+"7")
        for n in range(7,8+math.ceil(largest/24)):sheet.row_dimensions[n].height=18
        sheet.freeze_panes="A7"
    notes=workbook.create_sheet("검수범위")
    for key,value in findings.items():notes.append([key,json.dumps(value,ensure_ascii=False) if isinstance(value,(dict,list)) else str(value)])
    notes.column_dimensions["A"].width=32;notes.column_dimensions["B"].width=110
    for r in notes:
        for c in r:c.alignment=Alignment(wrap_text=True,vertical="top")
    workbook.save(BUNDLE/"comparison.xlsx")
    # Validate actual embedded images, IDs and referenced files, not just an XLSX extension.
    with zipfile.ZipFile(BUNDLE/"comparison.xlsx") as z:
        embedded=len([n for n in z.namelist() if n.startswith("xl/media/")])
    assert embedded==2*len(screens),(embedded,len(screens))
    for row in rows:
        for key in ("웹캡처","모바일캡처","정적시안"):
            assert (BUNDLE/row[key]).is_file(),(row["화면ID"],key)
    preserved=json.loads((OUT/"baseline.json").read_text())
    changed=[name for name,digest in preserved["files"].items() if not (ROOT/name).is_file() or hashlib.sha256((ROOT/name).read_bytes()).hexdigest()!=digest]
    assert not changed,changed[:10]
    import subprocess
    committed_changes=subprocess.check_output(["git","diff","--name-only",manifest["rollback_commit"]],cwd=ROOT,text=True).splitlines()
    assert not [name for name in committed_changes if name in preserved["files"]],committed_changes
    (BUNDLE/"source-preservation.json").write_text(json.dumps({"rollback_commit":preserved["rollback_commit"],"protected_files":len(preserved["files"]),"changed":[],"regression_sources_preserved":True},ensure_ascii=False,indent=2))
    # Never ship source code, credentials, the SDK bootstrap response, or scripts that initiate payment.
    for file in (BUNDLE/"screens").glob("*.html"):
        text=file.read_text()
        assert not re.search(r"(?:appkey|serviceKey|api_key)\s*=",text,re.I),file.name
    zip_path=OUT/"homenstay-all-screens-design.zip"
    with zipfile.ZipFile(zip_path,"w",zipfile.ZIP_DEFLATED,compresslevel=6) as z:
        for file in sorted(BUNDLE.rglob("*")):
            if file.is_file():z.write(file,file.relative_to(BUNDLE))
    print(json.dumps({"screens":len(screens),"captures":len(results["screens"]),"excel_images":embedded,"zip_bytes":zip_path.stat().st_size,"protected_files_unchanged":len(preserved["files"]),"needs_design_review":len(findings["design_needs"])},ensure_ascii=False))

if __name__=="__main__":
    main()
