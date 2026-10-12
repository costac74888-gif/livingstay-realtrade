"""Extend the real comparison workbook; archive new evidence without rewriting history."""
import hashlib
import json
import os
from datetime import datetime, timezone
from pathlib import Path
import zipfile
import openpyxl
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "generated-artifacts/hs2-rebuild"


def load(path):
    return json.loads(path.read_text())


def references(screen):
    view = screen["view"]
    if screen["role"] == "관리자":
        return [], "NOT_OBSERVED: 경쟁사 비공개 관리자 화면; 기존 자체 권한 경계 보존"
    if screen["role"] == "임대인":
        if "guide" in view:
            return ["33m2-host"], "공개 임대인 진입 소개만 관찰; 내부 관리 화면과 구분"
        return [], "NOT_OBSERVED: 방관리·계약·채팅 모두 로그인으로 이동; 내부 동등성 주장 금지"
    if any(word in view for word in ("detail", "gallery", "amenities", "price", "calendar")):
        return ["33m2-detail", "airbnb-detail"], "기간형 요금은 삼삼엠투, 숙박 사진·예약 요약 위계는 Airbnb 공개 상세 참고"
    if any(word in view for word in ("search", "filter", "map", "favorite")):
        return ["33m2-search", "airbnb-search"], "주·월 지도/목록은 삼삼엠투, 숙박 사진/목록은 Airbnb 공개 검색 참고"
    if any(word in view for word in ("login", "password", "contract", "chat", "booking", "payment", "account", "signup")):
        return [], "NOT_OBSERVED: 인증 후 내부 계약·채팅·결제는 공개 로그인 진입과 구분"
    return ["33m2-main", "airbnb-main"], "공개 탐색 진입 위계만 참고; 해당 내부 기능 직접 관찰 아님"


def main():
    delivery = OUT / "delivery"; delivery.mkdir(parents=True, exist_ok=True)
    registry = load(ROOT / "scripts/hs2_design_registry.json")
    preview = OUT / "preview-final-current"; capture = load(preview / "capture-report.json")
    snapshots = {(x["id"], x["width"]): x for x in capture["results"]}
    refs = []
    for folder in ("reference-current", "reference-detail"):
        for row in load(OUT / folder / "observations.json"):
            row["folder"] = folder; refs.append(row)
    baseline = load(OUT / "baseline/previous-evidence-manifest.json")
    changed = [r["path"] for r in baseline if not (ROOT / r["path"]).is_file()
               or hashlib.sha256((ROOT / r["path"]).read_bytes()).hexdigest() != r["sha256"]]
    runs = sorted((OUT / "regression-current-combined").glob("*/summary.json"), key=lambda p:p.stat().st_mtime)
    checks = load(runs[-1]) if runs else []
    fence = runs[-1].parent / "source-preservation.json" if runs else None
    preservation = load(fence) if fence and fence.exists() else {}
    regression = dict(checks=len(checks), passed=sum(r["exit_code"] == 0 for r in checks),
                      failed=[r["name"] for r in checks if r["exit_code"] != 0],
                      complete=bool(checks) and bool(preservation.get("unchanged")),
                      source_preserved=bool(preservation.get("unchanged")))
    summary = {
        "recorded_at": datetime.now(timezone.utc).isoformat(),
        "overall": "INCOMPLETE", "historical_evidence_unchanged": not changed,
        "changed_historical_files": changed, "visual_basic": capture["summary"],
        "native_mapped_items": sum(s["native"] for s in registry["screens"]),
        "local_presentation_items": sum(not s["native"] for s in registry["screens"]),
        "functional_acceptance": "NOT_APPROVED", "real_accounts_connected": False,
        "real_toss": "BLOCKED_DOCUMENT_EXAMPLE_KEYS", "production_modified": False,
        "provider_or_delivery_executed": False, "physical_device_verified": False,
        "owned_runtime_persists_after_restart": False,
        "latest_regression": regression,
    }
    functional = OUT / "functional-final/summary.json"
    if functional.exists():
        result = load(functional)
        summary["owned_functional_scenarios"] = {
            "status": result["status"], "checks": len(result["checks"]),
            "actual_toss": result["actual_toss"],
            "actual_money_executed": result["actual_money_executed"],
            "production_writes": result["production_writes"],
        }
    summary["desktop_marker_highlight"] = "UNVERIFIED_FAILED_CHECK"
    summary["earlier_whole_run"] = {"checks":164,"passed":158,"failed":6,
        "bank_standalone_recheck":"PASS","review_digest_rechecks":"4 PASS",
        "note":"Prior failures are retained; not replaced with targeted passes."}
    (delivery / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2))
    # Preserve original sheets AND their embedded historical screenshots.
    wb = openpyxl.load_workbook(ROOT / "generated-artifacts/hs2-integrated-preview/homenstay-running-preview-comparison.xlsx")
    for sheet in wb:
        sheet.title = ("이전검수_" + sheet.title)[:31]
    info = wb.create_sheet("최신 검수 요약", 0)
    for key, value in summary.items():
        info.append([key, json.dumps(value, ensure_ascii=False) if isinstance(value,(dict,list)) else str(value)])
    info.append(["주의", "화면 캡처·HTTP 200·가상 원장 PASS는 전체 기능/실계정/실제 토스 완료가 아닙니다."])
    info.append(["기존 자료", "이전검수 시트의 올리브 채택 설명은 과거 지시입니다. 최신 원본 관찰과 분리합니다."])
    info.column_dimensions["A"].width = 35; info.column_dimensions["B"].width = 105
    origin = "https://" + os.environ["REPLIT_DEV_DOMAIN"]
    thumbs = delivery / "thumbnails"; thumbs.mkdir(exist_ok=True)

    def picture(sheet, file, anchor, name, width=180, height=250):
        if not file.exists(): return
        target = thumbs / (name + ".jpg")
        with Image.open(file) as source:
            source.thumbnail((width, height)); source.convert("RGB").save(target, quality=65)
        image = ExcelImage(str(target)); sheet.add_image(image, anchor)

    for role in ("사용자","임대인","관리자"):
        sheet = wb.create_sheet("최신_" + role)
        sheet.append(["ID","화면·종류","현재 실행 URL","실제 기능 연결 범위","원본·선정 이유","관찰일·한계",
                      "360","390","430","1440","원본 삼삼엠투","원본 Airbnb"])
        sheet.freeze_panes="A2"; sheet.auto_filter.ref="A1:L1"
        for screen in (s for s in registry["screens"] if s["role"] == role):
            ids, reason = references(screen)
            observed = [r for r in refs if r["id"] in ids]
            source = "\n".join(sorted({r.get("final_url",r["url"]) for r in observed}))
            row = sheet.max_row + 1
            proof = [snapshots.get((screen["id"], w),{}) for w in (360,390,430,1440)]
            scope = ("API·가상 원장 연결(실계정/실제 결제 아님)" if screen["native"]
                     else "LOCAL_PRESENTATION_ONLY · 데이터 저장 기능 완료 아님")
            sheet.append([screen["id"],screen["title"]+" / "+screen["kind"],origin+screen["preview_path"],
                scope,reason+"\n"+source,
                "\n".join(sorted({r["observed_at"] for r in observed}))+"\n시각적 원본 일치 승인 아님",
                *["기초 검사 PASS" if p.get("visual_basic_pass") else "FAIL/미캡처" for p in proof],
                "공개 원본만" if any(i.startswith("33m2") for i in ids) else "NOT_OBSERVED",
                "공개 원본만" if any(i.startswith("airbnb") for i in ids) else "NOT_OBSERVED"])
            sheet.row_dimensions[row].height = 215
            for col, p in zip("GHIJ", proof):
                if p.get("capture"):
                    picture(sheet,preview/p["capture"],col+str(row),screen["id"]+"-"+str(p["width"]))
            for prefix,col in (("33m2","K"),("airbnb","L")):
                ref = next((r for r in observed if r["id"].startswith(prefix) and r["width"]==390 and r.get("capture")),None)
                if ref:
                    picture(sheet,OUT/ref["folder"]/ref["capture"],col+str(row),screen["id"]+"-"+prefix)
            for cell in sheet[row]:
                cell.alignment=Alignment(vertical="top",wrap_text=True)
            sheet.cell(row,3).hyperlink=origin+screen["preview_path"]
        for col in "ABCDEF": sheet.column_dimensions[col].width = 30 if col!="E" else 55
        for col in "GHIJKL": sheet.column_dimensions[col].width = 27
        for cell in sheet[1]:
            cell.font=Font(bold=True,color="FFFFFF");cell.fill=PatternFill("solid",fgColor="2563EB")
        sheet.auto_filter.ref=f"A1:L{sheet.max_row}"
    workbook = delivery / "homenstay-rebuild-comparison.xlsx"; wb.save(workbook)
    text = ("HOME & STAY 개발 재구축 검수 자료\n\n"
            "이전 검수 자료는 원위치에서 보존했고, Excel 이전검수 시트도 유지했습니다.\n"
            "최신 캡처의 PASS는 가로 넘침/이미지/HTTP 기초 확인일 뿐 기능 또는 디자인 동등성 승인 아닙니다.\n"
             "실제 계정 연결 및 실제 토스 테스트는 완료되지 않았습니다. 운영 DB/게시/실제 금전/발송은 변경하지 않았습니다.\n"
             "가상 예약·채팅은 현재 PostgreSQL 실행 중 새로고침 뒤에도 유지되지만, Preview 재시작 뒤 영구 보존은 완료되지 않았습니다.\n"
            "최신 전체 회귀: "+json.dumps(regression,ensure_ascii=False)+"\n")
    (delivery/"README.txt").write_text(text)
    archive = delivery / "homenstay-rebuild-captures.zip"
    with zipfile.ZipFile(archive,"w",zipfile.ZIP_DEFLATED,compresslevel=3) as z:
        for folder in (preview,OUT/"reference-current",OUT/"reference-detail",OUT/"functional-final"):
            for file in folder.rglob("*"):
                if file.is_file(): z.write(file,file.relative_to(OUT))
        for file in (delivery/"README.txt",delivery/"summary.json",OUT/"toss-key-metadata.json"):
            z.write(file,file.relative_to(OUT))
        if runs:
            z.write(runs[-1],"regression/summary.json")
            if fence and fence.exists():
                z.write(fence,"regression/source-preservation.json")
        z.write(OUT/"baseline/previous-evidence-manifest.json","baseline/previous-evidence-manifest.json")
        for name in ("current-map.log","current-map-recheck.log","current-map-final.log",
                     "current-map-confirmed.log","current-map-desktop.log"):
            file=OUT/name
            if file.exists(): z.write(file,"unresolved-map-check/"+name)
    print(json.dumps(summary,ensure_ascii=False));print(workbook.relative_to(ROOT));print(archive.relative_to(ROOT))


if __name__ == "__main__":
    main()
