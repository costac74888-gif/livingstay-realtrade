"""Build the revised relay operations manual without changing app settings."""
from pathlib import Path
import re
from zipfile import ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "docs/공공_API_중계서버_구축운영_매뉴얼_v2_2026-10-06.md"
OUTPUT = ROOT / "exports/public_api_relay_manual_v2_2026-10-06.docx"


def font(run, name="맑은 고딕", size=10.5, bold=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    props = run._element.get_or_add_rPr()
    fonts = props.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        props.insert(0, fonts)
    fonts.set(qn("w:eastAsia"), "맑은 고딕" if name == "Consolas" else name)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def shade(element, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    element.append(shd)


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(1.9), Cm(1.9)
    section.left_margin, section.right_margin = Cm(1.8), Cm(1.8)
    section.header_distance = section.footer_distance = Cm(0.8)
    section.different_first_page_header_footer = True
    normal = doc.styles["Normal"]
    normal.font.name = "맑은 고딕"
    normal.font.size = Pt(10.5)
    normal.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(6)
    for name, size in [("Title", 25), ("Heading 1", 17), ("Heading 2", 13)]:
        style = doc.styles[name]
        style.font.name = "맑은 고딕"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string("244A5B")
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "맑은 고딕")
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(14)
        style.paragraph_format.space_after = Pt(7)
    header = section.header.paragraphs[0]
    font(header.add_run("공공 API 중계서버 구축·운영 매뉴얼  |  v2"), size=9, color="657580")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    font(footer.add_run("2026-10-06  ·  v2  |  "), size=9, color="657580")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.core_properties.title = "공공 API 중계서버 구축·운영 매뉴얼 v2"
    doc.core_properties.subject = "원본 오류 수정, 단계별 활성화 및 보안·운영 절차"
    doc.core_properties.author = "홈앤스테이"
    doc.core_properties.version = "2.0"

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("```"):
            code = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                code.append(lines[i])
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.25)
            p.paragraph_format.right_indent = Cm(0.25)
            p.paragraph_format.line_spacing = 1.05
            p.paragraph_format.keep_together = True
            shade(p._p.get_or_add_pPr(), "F0F4F6")
            font(p.add_run("\n".join(code)), name="Consolas", size=9.5, color="203746")
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                values = [x.strip() for x in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", x) for x in values):
                    rows.append(values)
                i += 1
            table = doc.add_table(rows=1, cols=len(rows[0]))
            table.style = "Table Grid"
            table.autofit = False
            widths = [5.0, 12.4] if len(rows[0]) == 2 else (
                [3.6, 6.6, 7.2] if len(rows[0]) == 3 else [4.1, 1.4, 4.8, 7.1]
            )
            for col, width in zip(table.columns, widths):
                col.width = Cm(width)
            for n, values in enumerate(rows):
                cells = table.rows[0].cells if n == 0 else table.add_row().cells
                props = table.rows[n]._tr.get_or_add_trPr()
                props.append(OxmlElement("w:cantSplit"))
                if n == 0:
                    props.append(OxmlElement("w:tblHeader"))
                for j, (cell, value) in enumerate(zip(cells, values)):
                    cell.width = Cm(widths[j])
                    if n == 0:
                        shade(cell._tc.get_or_add_tcPr(), "E6EEF2")
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_after = Pt(4)
                    p.paragraph_format.space_before = Pt(4)
                    p.paragraph_format.line_spacing = 1.12
                    font(p.add_run(value), size=9, bold=n == 0)
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
            continue
        elif line.startswith("# "):
            doc.add_paragraph(line[2:], style="Title")
        elif line.startswith("## "):
            p = doc.add_paragraph(line[3:], style="Heading 1")
            if line.startswith("## 1. "):
                p.paragraph_format.page_break_before = True
        elif line.startswith("### "):
            doc.add_paragraph(line[4:], style="Heading 2")
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            font(p.add_run(line[2:]))
        elif re.match(r"^\d+\.\s", line):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.35)
            p.paragraph_format.first_line_indent = Cm(-0.35)
            font(p.add_run(line))
        else:
            p = doc.add_paragraph()
            font(p.add_run(line))
        i += 1

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    reopened = Document(OUTPUT)
    with ZipFile(OUTPUT) as z:
        assert z.testzip() is None
    paragraphs = [p.text for p in reopened.paragraphs]
    headings = [p.text for p in reopened.paragraphs if p.style.name == "Heading 1"]
    assert all(any(p.startswith(f"{n}. ") for p in headings) for n in range(1, 16))
    text = "\n".join(paragraphs) + "\n" + "\n".join(
        c.text for t in reopened.tables for r in t.rows for c in r.cells
    )
    for phrase in ["재게시", "기존 호출자", "신뢰", "umask 027", "Autoscale", "Scheduled", "경로 부분", "v2"]:
        assert phrase in text, phrase
    print(f"Created: {OUTPUT.relative_to(ROOT)}")
    print(f"Validation PASS: zip integrity, Word reopen, 15 sections, {len(reopened.tables)} tables, correction coverage")


if __name__ == "__main__":
    build()
