"""2024 호텔업 운영현황 지역별 XLSX ZIP을 중복 없이 적재한다."""

import argparse
import hashlib
import io
import os
import re
import zipfile

from openpyxl import load_workbook
from psycopg2.extras import execute_values

from address_utils import sido_core
SOURCE_NAME = "한국호텔업협회 호텔업 운영현황"
DEFAULT_YEAR = 2024
EXPECTED_REGION_FILE_COUNT = 16
EXPECTED_SIDOS = frozenset({
    "서울", "부산", "대구", "인천", "광주", "대전", "울산",
    "경기", "강원", "충청북", "충청남", "전라북", "전라남",
    "경상북", "경상남", "제주",
})


def _number(value, integer=False):
    if value is None or value == "":
        return None
    try:
        number = float(str(value).replace(",", "").strip())
        return int(number) if integer else number
    except (TypeError, ValueError):
        return None


def parse_operation_zip(zip_path, reference_year=DEFAULT_YEAR):
    """지역별 연간 시트의 시군구 전체행과 등급행을 원본 순서로 반환한다."""
    with open(zip_path, "rb") as source:
        archive_bytes = source.read()
    source_sha256 = hashlib.sha256(archive_bytes).hexdigest()
    records = []
    with zipfile.ZipFile(io.BytesIO(archive_bytes)) as archive:
        names = sorted(
            name for name in archive.namelist()
            if name.endswith("지역 원데이터.xlsx") and "/전국/" not in name
        )
        for name in names:
            path_parts = name.split("/")
            sido_name = sido_core(path_parts[-2])
            workbook = load_workbook(
                io.BytesIO(archive.read(name)), read_only=True, data_only=True
            )
            worksheet = workbook["Sheet1"]
            current_sgg = None
            for row_number, row in enumerate(
                worksheet.iter_rows(min_row=2, values_only=True), start=2
            ):
                label = str(row[0] or "")
                grade = str(row[16] or "").strip()
                stripped = re.sub(r"\s+", "", label)
                if not stripped or not grade:
                    continue
                if "/전체" in stripped:
                    current_sgg = None
                elif label[:1].isspace():
                    # 들여쓴 등급행은 직전 시군구(또는 시도 합계)에 속한다.
                    pass
                else:
                    current_sgg = stripped
                records.append((
                    name, row_number, sido_name, current_sgg, grade,
                    _number(row[8]), _number(row[9]), _number(row[10]),
                    _number(row[11]), _number(row[12], True),
                    _number(row[13], True), _number(row[14], True),
                    _number(row[15], True),
                ))
            workbook.close()
    return source_sha256, records


def validate_operation_records(records):
    source_files = {row[0] for row in records}
    sidos = {row[2] for row in records}
    sgg_overall = {
        row[2] for row in records if row[3] is not None and row[4] == "전체"
    }
    if (
        len(source_files) != EXPECTED_REGION_FILE_COUNT
        or sidos != EXPECTED_SIDOS
        or sgg_overall != sidos
    ):
        raise ValueError(
            "지역별 원본이 완전하지 않습니다: "
            f"파일 {len(source_files)}개, 시도 {len(sidos)}개, "
            f"시군구 전체행 보유 시도 {len(sgg_overall)}개, "
            f"누락 {sorted(EXPECTED_SIDOS - sidos)}, "
            f"예상 외 {sorted(sidos - EXPECTED_SIDOS)}"
        )


def import_operation_zip(zip_path, reference_year=DEFAULT_YEAR, conn=None):
    from db import get_conn

    source_sha256, records = parse_operation_zip(zip_path, reference_year)
    validate_operation_records(records)
    owned_connection = conn is None
    conn = conn or get_conn()
    cur = conn.cursor()
    try:
        cur.execute("""
            INSERT INTO hotel_operation_source_versions
                (reference_year, source_name, source_file, source_sha256)
            VALUES (%s, %s, %s, %s)
            ON CONFLICT (source_sha256) DO NOTHING
            RETURNING id
        """, (reference_year, SOURCE_NAME, os.path.basename(zip_path), source_sha256))
        inserted = cur.fetchone()
        if not inserted:
            conn.rollback()
            return {"inserted": False, "rows": 0, "sha256": source_sha256}
        version_id = inserted["id"]
        execute_values(cur, """
            INSERT INTO hotel_operation_metrics
                (version_id, source_file, source_row_number, sido_name, sgg_name,
                 grade, occupancy_rate, adr, revpar, foreign_guest_rate,
                 available_room_nights, sold_room_nights, room_count, business_count)
            VALUES %s
        """, [(version_id, *record) for record in records])
        conn.commit()
        return {"inserted": True, "rows": len(records), "sha256": source_sha256}
    except Exception:
        conn.rollback()
        raise
    finally:
        cur.close()
        if owned_connection:
            conn.close()


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("zip_path")
    parser.add_argument("--year", type=int, default=DEFAULT_YEAR)
    args = parser.parse_args()
    result = import_operation_zip(args.zip_path, args.year)
    print(("적재 완료" if result["inserted"] else "이미 적재된 원본"), result["rows"])


if __name__ == "__main__":
    main()